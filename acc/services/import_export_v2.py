"""
Advanced Import/Export System V2
نظام متقدم للاستيراد والتصدير بأعلى مستوى من الأداء والموثوقية
"""
import pandas as pd
import chardet
import json
import io
import os
from datetime import datetime
from typing import List, Dict, Any, Tuple, Optional
from flask import Response, make_response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import csv


class UniversalImportExport:
    """نظام موحد ومتقدم للاستيراد والتصدير"""
    
    # تعريف الحقول المدعومة
    FIELD_MAPPINGS = {
        # Arabic
        'الكود': 'code',
        'الاسم': 'name',
        'الهاتف': 'phone',
        'البريد الإلكتروني': 'email',
        'البريد الالكتروني': 'email',
        'العنوان': 'address',
        'الرقم القومي': 'national_id',
        'الحالة': 'status',
        'الملاحظات': 'notes',
        'ملاحظات': 'notes',
        'تاريخ التسجيل': 'created_at',
        'تاريخ الإنشاء': 'created_at',
        
        # English
        'code': 'code',
        'name': 'name',
        'phone': 'phone',
        'email': 'email',
        'address': 'address',
        'national_id': 'national_id',
        'national id': 'national_id',
        'status': 'status',
        'notes': 'notes',
        'created_at': 'created_at',
        'created at': 'created_at',
        'date': 'created_at',
    }
    
    # القيم الافتراضية للحقول
    FIELD_DEFAULTS = {
        'status': 'نشط',
        'phone': None,
        'email': None,
        'address': None,
        'national_id': None,
        'notes': None
    }
    
    # قواعد التحقق من صحة البيانات
    VALIDATION_RULES = {
        'name': {
            'required': True,
            'max_length': 100,
            'min_length': 2,
            'pattern': None
        },
        'phone': {
            'required': False,
            'max_length': 20,
            'pattern': r'^[\d\s\-\+\(\)]+$',
            'message': 'رقم الهاتف يجب أن يحتوي على أرقام فقط'
        },
        'email': {
            'required': False,
            'pattern': r'^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$',
            'message': 'البريد الإلكتروني غير صحيح'
        },
        'national_id': {
            'required': False,
            'max_length': 20,
            'pattern': r'^[\d]+$',
            'message': 'الرقم القومي يجب أن يحتوي على أرقام فقط'
        },
        'status': {
            'required': False,
            'choices': ['نشط', 'غير نشط', 'active', 'inactive'],
            'default': 'نشط'
        }
    }
    
    @classmethod
    def detect_encoding(cls, file_content: bytes) -> str:
        """اكتشاف ترميز الملف بدقة عالية"""
        # محاولة اكتشاف الترميز باستخدام chardet
        result = chardet.detect(file_content)
        encoding = result['encoding']
        confidence = result['confidence']
        
        # إذا كانت الثقة منخفضة، جرب ترميزات عربية شائعة
        if confidence < 0.7:
            encodings = ['utf-8-sig', 'utf-8', 'windows-1256', 'iso-8859-6']
            for enc in encodings:
                try:
                    file_content.decode(enc)
                    return enc
                except:
                    continue
        
        return encoding or 'utf-8'
    
    @classmethod
    def validate_field(cls, field_name: str, value: Any, row_num: int) -> Tuple[Any, Optional[str]]:
        """التحقق من صحة قيمة حقل واحد"""
        if field_name not in cls.VALIDATION_RULES:
            return value, None
        
        rules = cls.VALIDATION_RULES[field_name]
        
        # تحقق من الحقل المطلوب
        if rules.get('required') and not value:
            return None, f"السطر {row_num}: {field_name} مطلوب"
        
        # إذا كانت القيمة فارغة وليست مطلوبة
        if not value and not rules.get('required'):
            return cls.FIELD_DEFAULTS.get(field_name), None
        
        # تحويل القيمة إلى نص
        value = str(value).strip()
        
        # تحقق من الطول الأقصى
        if 'max_length' in rules and len(value) > rules['max_length']:
            return None, f"السطر {row_num}: {field_name} طويل جداً (الحد الأقصى {rules['max_length']} حرف)"
        
        # تحقق من الطول الأدنى
        if 'min_length' in rules and len(value) < rules['min_length']:
            return None, f"السطر {row_num}: {field_name} قصير جداً (الحد الأدنى {rules['min_length']} حرف)"
        
        # تحقق من النمط
        if 'pattern' in rules and rules['pattern']:
            import re
            if not re.match(rules['pattern'], value):
                return None, f"السطر {row_num}: {rules.get('message', f'{field_name} غير صحيح')}"
        
        # تحقق من الخيارات المحددة
        if 'choices' in rules:
            if value.lower() not in [c.lower() for c in rules['choices']]:
                # حاول تصحيح القيمة
                if value.lower() in ['active', 'نشط']:
                    value = 'نشط'
                elif value.lower() in ['inactive', 'غير نشط']:
                    value = 'غير نشط'
                else:
                    value = rules.get('default', rules['choices'][0])
        
        return value, None
    
    @classmethod
    def import_csv(cls, file_content: bytes) -> Tuple[pd.DataFrame, List[str], Dict[str, Any]]:
        """استيراد ملف CSV مع معالجة متقدمة"""
        errors = []
        warnings = []
        metadata = {
            'total_rows': 0,
            'valid_rows': 0,
            'encoding': None,
            'delimiter': None
        }
        
        try:
            # اكتشاف الترميز
            encoding = cls.detect_encoding(file_content)
            metadata['encoding'] = encoding
            
            # محاولة قراءة الملف
            text = file_content.decode(encoding)
            
            # اكتشاف الفاصل
            sample = text[:1024]
            delimiter = ','
            for delim in [',', ';', '\t', '|']:
                if delim in sample:
                    delimiter = delim
                    break
            metadata['delimiter'] = delimiter
            
            # قراءة البيانات باستخدام pandas
            df = pd.read_csv(io.StringIO(text), delimiter=delimiter)
            metadata['total_rows'] = len(df)
            
            # تنظيف أسماء الأعمدة
            df.columns = [col.strip() for col in df.columns]
            
            # تحويل أسماء الأعمدة
            column_mapping = {}
            for col in df.columns:
                normalized = cls.FIELD_MAPPINGS.get(col, col.lower().replace(' ', '_'))
                column_mapping[col] = normalized
            
            df = df.rename(columns=column_mapping)
            
            # التحقق من وجود عمود الاسم
            if 'name' not in df.columns:
                errors.append("لم يتم العثور على عمود 'الاسم' أو 'name' في الملف")
                return pd.DataFrame(), errors, metadata
            
            # تنظيف البيانات
            df = df.fillna('')
            df = df.applymap(lambda x: str(x).strip() if pd.notna(x) else '')
            
            # التحقق من صحة البيانات
            valid_rows = []
            for idx, row in df.iterrows():
                row_errors = []
                validated_row = {}
                
                for field in ['name', 'phone', 'email', 'national_id', 'address', 'status', 'notes']:
                    value = row.get(field, '')
                    validated_value, error = cls.validate_field(field, value, idx + 2)
                    
                    if error:
                        row_errors.append(error)
                    else:
                        validated_row[field] = validated_value
                
                if not row_errors:
                    valid_rows.append(validated_row)
                else:
                    errors.extend(row_errors)
            
            # إنشاء DataFrame من الصفوف الصحيحة
            result_df = pd.DataFrame(valid_rows)
            metadata['valid_rows'] = len(result_df)
            
            return result_df, errors, metadata
            
        except Exception as e:
            errors.append(f"خطأ في قراءة ملف CSV: {str(e)}")
            return pd.DataFrame(), errors, metadata
    
    @classmethod
    def import_excel(cls, file_content: bytes) -> Tuple[pd.DataFrame, List[str], Dict[str, Any]]:
        """استيراد ملف Excel مع معالجة متقدمة"""
        errors = []
        metadata = {
            'total_rows': 0,
            'valid_rows': 0,
            'sheets': []
        }
        
        try:
            # قراءة ملف Excel
            excel_file = pd.ExcelFile(io.BytesIO(file_content))
            metadata['sheets'] = excel_file.sheet_names
            
            # قراءة أول ورقة
            df = pd.read_excel(excel_file, sheet_name=0)
            metadata['total_rows'] = len(df)
            
            # تنظيف أسماء الأعمدة
            df.columns = [str(col).strip() for col in df.columns]
            
            # تحويل أسماء الأعمدة
            column_mapping = {}
            for col in df.columns:
                normalized = cls.FIELD_MAPPINGS.get(col, col.lower().replace(' ', '_'))
                column_mapping[col] = normalized
            
            df = df.rename(columns=column_mapping)
            
            # التحقق من وجود عمود الاسم
            if 'name' not in df.columns:
                errors.append("لم يتم العثور على عمود 'الاسم' أو 'name' في الملف")
                return pd.DataFrame(), errors, metadata
            
            # تنظيف البيانات
            df = df.fillna('')
            df = df.applymap(lambda x: str(x).strip() if pd.notna(x) else '')
            
            # التحقق من صحة البيانات
            valid_rows = []
            for idx, row in df.iterrows():
                row_errors = []
                validated_row = {}
                
                for field in ['name', 'phone', 'email', 'national_id', 'address', 'status', 'notes']:
                    value = row.get(field, '')
                    validated_value, error = cls.validate_field(field, value, idx + 2)
                    
                    if error:
                        row_errors.append(error)
                    else:
                        validated_row[field] = validated_value
                
                if not row_errors:
                    valid_rows.append(validated_row)
                else:
                    errors.extend(row_errors)
            
            # إنشاء DataFrame من الصفوف الصحيحة
            result_df = pd.DataFrame(valid_rows)
            metadata['valid_rows'] = len(result_df)
            
            return result_df, errors, metadata
            
        except Exception as e:
            errors.append(f"خطأ في قراءة ملف Excel: {str(e)}")
            return pd.DataFrame(), errors, metadata
    
    @classmethod
    def import_json(cls, file_content: bytes) -> Tuple[pd.DataFrame, List[str], Dict[str, Any]]:
        """استيراد ملف JSON مع معالجة متقدمة"""
        errors = []
        metadata = {
            'total_rows': 0,
            'valid_rows': 0,
            'encoding': None
        }
        
        try:
            # اكتشاف الترميز
            encoding = cls.detect_encoding(file_content)
            metadata['encoding'] = encoding
            
            # قراءة JSON
            data = json.loads(file_content.decode(encoding))
            
            # التعامل مع تنسيقات JSON المختلفة
            if isinstance(data, list):
                items = data
            elif isinstance(data, dict):
                # محاولة إيجاد قائمة البيانات
                for key in ['data', 'items', 'records', 'customers', 'العملاء']:
                    if key in data and isinstance(data[key], list):
                        items = data[key]
                        break
                else:
                    items = [data]
            else:
                errors.append("تنسيق JSON غير صحيح")
                return pd.DataFrame(), errors, metadata
            
            metadata['total_rows'] = len(items)
            
            # معالجة البيانات
            valid_rows = []
            for idx, item in enumerate(items):
                if not isinstance(item, dict):
                    errors.append(f"العنصر {idx + 1}: ليس كائن JSON صحيح")
                    continue
                
                # تحويل أسماء الحقول
                normalized_item = {}
                for key, value in item.items():
                    normalized_key = cls.FIELD_MAPPINGS.get(key, key.lower().replace(' ', '_'))
                    normalized_item[normalized_key] = value
                
                # التحقق من صحة البيانات
                row_errors = []
                validated_row = {}
                
                for field in ['name', 'phone', 'email', 'national_id', 'address', 'status', 'notes']:
                    value = normalized_item.get(field, '')
                    validated_value, error = cls.validate_field(field, value, idx + 1)
                    
                    if error:
                        row_errors.append(error)
                    else:
                        validated_row[field] = validated_value
                
                if not row_errors:
                    valid_rows.append(validated_row)
                else:
                    errors.extend(row_errors)
            
            # إنشاء DataFrame
            result_df = pd.DataFrame(valid_rows)
            metadata['valid_rows'] = len(result_df)
            
            return result_df, errors, metadata
            
        except json.JSONDecodeError as e:
            errors.append(f"خطأ في تحليل JSON: {str(e)}")
            return pd.DataFrame(), errors, metadata
        except Exception as e:
            errors.append(f"خطأ في قراءة ملف JSON: {str(e)}")
            return pd.DataFrame(), errors, metadata
    
    @classmethod
    def import_file(cls, file_content: bytes, filename: str) -> Tuple[pd.DataFrame, List[str], Dict[str, Any]]:
        """استيراد ملف بناءً على نوعه"""
        filename_lower = filename.lower()
        
        if filename_lower.endswith('.csv'):
            return cls.import_csv(file_content)
        elif filename_lower.endswith(('.xlsx', '.xls')):
            return cls.import_excel(file_content)
        elif filename_lower.endswith('.json'):
            return cls.import_json(file_content)
        else:
            return pd.DataFrame(), ["نوع الملف غير مدعوم. الرجاء استخدام CSV, Excel, أو JSON."], {}
    
    @classmethod
    def export_csv(cls, data: List[Dict], filename: str = "export") -> Response:
        """تصدير البيانات إلى CSV مع تنسيق احترافي"""
        output = io.StringIO()
        
        # الحقول المراد تصديرها
        fields = ['code', 'name', 'phone', 'email', 'national_id', 'address', 'status', 'notes', 'created_at']
        headers = ['الكود', 'الاسم', 'الهاتف', 'البريد الإلكتروني', 'الرقم القومي', 'العنوان', 'الحالة', 'الملاحظات', 'تاريخ التسجيل']
        
        writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
        
        # كتابة معلومات التصدير
        writer.writerow([f'تم التصدير بتاريخ: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}'])
        writer.writerow([f'عدد السجلات: {len(data)}'])
        writer.writerow([])  # سطر فارغ
        
        # كتابة العناوين
        writer.writerow(headers)
        
        # كتابة البيانات
        for item in data:
            row = []
            for field in fields:
                value = item.get(field, '')
                if field == 'created_at' and value:
                    # تنسيق التاريخ
                    if hasattr(value, 'strftime'):
                        value = value.strftime('%Y-%m-%d %H:%M:%S')
                row.append(value or '')
            writer.writerow(row)
        
        # إنشاء الاستجابة مع BOM لدعم العربية في Excel
        response = make_response('\ufeff' + output.getvalue())
        response.headers['Content-Type'] = 'text/csv; charset=utf-8-sig'
        response.headers['Content-Disposition'] = f'attachment; filename={filename}_{datetime.now().strftime("%Y%m%d_%H%M%S")}.csv'
        
        return response
    
    @classmethod
    def export_excel(cls, data: List[Dict], filename: str = "export") -> Response:
        """تصدير البيانات إلى Excel مع تنسيق احترافي"""
        wb = Workbook()
        ws = wb.active
        ws.title = "البيانات"
        
        # إعداد التنسيق
        header_font = Font(bold=True, color="FFFFFF", size=12)
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        
        border = Border(
            left=Side(border_style="thin"),
            right=Side(border_style="thin"),
            top=Side(border_style="thin"),
            bottom=Side(border_style="thin")
        )
        
        # معلومات التصدير
        ws['A1'] = f'تم التصدير بتاريخ: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}'
        ws['A2'] = f'عدد السجلات: {len(data)}'
        ws.merge_cells('A1:I1')
        ws.merge_cells('A2:I2')
        
        # العناوين
        headers = ['الكود', 'الاسم', 'الهاتف', 'البريد الإلكتروني', 'الرقم القومي', 'العنوان', 'الحالة', 'الملاحظات', 'تاريخ التسجيل']
        fields = ['code', 'name', 'phone', 'email', 'national_id', 'address', 'status', 'notes', 'created_at']
        
        # كتابة العناوين
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border
        
        # كتابة البيانات
        for row_idx, item in enumerate(data, 5):
            for col_idx, field in enumerate(fields, 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                value = item.get(field, '')
                
                if field == 'created_at' and value:
                    if hasattr(value, 'strftime'):
                        value = value.strftime('%Y-%m-%d %H:%M:%S')
                
                cell.value = value or ''
                cell.border = border
                cell.alignment = Alignment(horizontal="right" if col_idx > 1 else "center")
        
        # ضبط عرض الأعمدة
        column_widths = [15, 25, 15, 25, 20, 30, 10, 30, 20]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        
        # إضافة ورقة إحصائيات
        ws_stats = wb.create_sheet(title="إحصائيات")
        
        # حساب الإحصائيات
        total_count = len(data)
        active_count = len([d for d in data if d.get('status') == 'نشط'])
        inactive_count = total_count - active_count
        
        stats_data = [
            ['الإحصائية', 'القيمة'],
            ['إجمالي العملاء', total_count],
            ['العملاء النشطون', active_count],
            ['العملاء غير النشطين', inactive_count],
            ['نسبة النشطين', f'{(active_count/total_count*100 if total_count > 0 else 0):.1f}%']
        ]
        
        for row_idx, row_data in enumerate(stats_data, 1):
            for col_idx, value in enumerate(row_data, 1):
                cell = ws_stats.cell(row=row_idx, column=col_idx)
                cell.value = value
                cell.border = border
                if row_idx == 1:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = header_alignment
        
        # ضبط عرض الأعمدة في ورقة الإحصائيات
        ws_stats.column_dimensions['A'].width = 20
        ws_stats.column_dimensions['B'].width = 15
        
        # حفظ الملف
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = make_response(output.getvalue())
        response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        response.headers['Content-Disposition'] = f'attachment; filename={filename}_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        
        return response
    
    @classmethod
    def export_json(cls, data: List[Dict], filename: str = "export") -> Response:
        """تصدير البيانات إلى JSON مع تنسيق احترافي"""
        # إعداد البيانات للتصدير
        export_data = {
            'metadata': {
                'exported_at': datetime.now().isoformat(),
                'total_records': len(data),
                'version': '2.0',
                'fields': ['code', 'name', 'phone', 'email', 'national_id', 'address', 'status', 'notes', 'created_at']
            },
            'data': []
        }
        
        # معالجة البيانات
        for item in data:
            record = {}
            for field in export_data['metadata']['fields']:
                value = item.get(field)
                if field == 'created_at' and value and hasattr(value, 'isoformat'):
                    value = value.isoformat()
                record[field] = value
            export_data['data'].append(record)
        
        # إنشاء الاستجابة
        response = make_response(json.dumps(export_data, ensure_ascii=False, indent=2))
        response.headers['Content-Type'] = 'application/json; charset=utf-8'
        response.headers['Content-Disposition'] = f'attachment; filename={filename}_{datetime.now().strftime("%Y%m%d_%H%M%S")}.json'
        
        return response


class ImportResult:
    """نتيجة عملية الاستيراد"""
    def __init__(self):
        self.success = False
        self.imported_count = 0
        self.skipped_count = 0
        self.error_count = 0
        self.errors = []
        self.warnings = []
        self.skipped_items = []
        self.preview_data = None
        self.metadata = {}
    
    def to_dict(self):
        return {
            'success': self.success,
            'imported_count': self.imported_count,
            'skipped_count': self.skipped_count,
            'error_count': self.error_count,
            'errors': self.errors[:10],  # أول 10 أخطاء فقط
            'warnings': self.warnings,
            'has_more_errors': len(self.errors) > 10,
            'total_errors': len(self.errors),
            'metadata': self.metadata
        }