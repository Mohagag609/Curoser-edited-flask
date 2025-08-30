"""
محرك التصدير العام
"""
import csv
import io
from typing import List, Dict, Any, Optional
from flask import g, make_response
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from .schemas import get_schema, ResourceSchema

class GenericExporter:
    """مصدر عام للبيانات"""
    
    def __init__(self, resource_name: str):
        self.resource_name = resource_name
        self.schema = get_schema(resource_name)
        if not self.schema:
            raise ValueError(f"مورد غير معرف: {resource_name}")
        
        self.model_class = self.schema.model_class
        self.current_project_id = g.current_project.id if hasattr(g, 'current_project') and g.current_project else None
    
    def export_csv(self) -> Any:
        """تصدير كـ CSV"""
        # جلب البيانات
        data = self._fetch_data()
        
        # إنشاء CSV
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=self._get_export_columns(), extrasaction='ignore')
        
        # كتابة العناوين
        headers = {}
        for col in self.schema.columns:
            headers[col['name']] = col['label']
        writer.writerow(headers)
        
        # كتابة البيانات
        for row in data:
            writer.writerow(row)
        
        # إنشاء الاستجابة
        response = make_response(output.getvalue())
        response.headers['Content-Type'] = 'text/csv; charset=utf-8-sig'
        response.headers['Content-Disposition'] = f'attachment; filename={self.resource_name}_export.csv'
        
        # إضافة BOM لدعم Excel العربي
        response.data = '\ufeff' + response.data
        
        return response
    
    def export_excel(self) -> Any:
        """تصدير كـ Excel"""
        # جلب البيانات
        data = self._fetch_data()
        
        # إنشاء Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = self.schema.plural_name
        
        # تنسيق العناوين
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        
        # كتابة العناوين
        for idx, col in enumerate(self.schema.columns, 1):
            cell = ws.cell(row=1, column=idx, value=col['label'])
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            
            # ضبط عرض العمود
            ws.column_dimensions[cell.column_letter].width = 20
        
        # كتابة البيانات
        for row_idx, row_data in enumerate(data, 2):
            for col_idx, col in enumerate(self.schema.columns, 1):
                value = row_data.get(col['name'], '')
                ws.cell(row=row_idx, column=col_idx, value=value)
        
        # حفظ في الذاكرة
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        # إنشاء الاستجابة
        response = make_response(output.getvalue())
        response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        response.headers['Content-Disposition'] = f'attachment; filename={self.resource_name}_export.xlsx'
        
        return response
    
    def generate_template(self, format: str = 'csv') -> Any:
        """توليد قالب للاستيراد"""
        # بيانات أمثلة
        example_data = []
        for i in range(3):
            row = {}
            for col in self.schema.columns:
                row[col['name']] = col.get('example', '')
            example_data.append(row)
        
        if format == 'xlsx':
            return self._generate_excel_template(example_data)
        else:
            return self._generate_csv_template(example_data)
    
    def _generate_csv_template(self, example_data: List[Dict]) -> Any:
        """توليد قالب CSV"""
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=self._get_export_columns())
        
        # كتابة العناوين
        headers = {}
        for col in self.schema.columns:
            headers[col['name']] = col['label']
        writer.writerow(headers)
        
        # كتابة الأمثلة
        writer.writerows(example_data)
        
        # إنشاء الاستجابة
        response = make_response(output.getvalue())
        response.headers['Content-Type'] = 'text/csv; charset=utf-8-sig'
        response.headers['Content-Disposition'] = f'attachment; filename={self.resource_name}_template.csv'
        
        # إضافة BOM
        response.data = '\ufeff' + response.data
        
        return response
    
    def _generate_excel_template(self, example_data: List[Dict]) -> Any:
        """توليد قالب Excel"""
        wb = Workbook()
        ws = wb.active
        ws.title = self.schema.plural_name
        
        # تنسيقات
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        required_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
        note_font = Font(italic=True, color="7F7F7F")
        
        # كتابة العناوين
        for idx, col in enumerate(self.schema.columns, 1):
            # العنوان
            cell = ws.cell(row=1, column=idx, value=col['label'])
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # علامة الحقول المطلوبة
            if col.get('required', False):
                cell.fill = required_fill
            
            # عرض العمود
            ws.column_dimensions[cell.column_letter].width = 20
        
        # إضافة ملاحظات
        notes_row = 2
        for idx, col in enumerate(self.schema.columns, 1):
            if 'note' in col:
                cell = ws.cell(row=notes_row, column=idx, value=col['note'])
                cell.font = note_font
        
        # كتابة الأمثلة
        start_row = 3 if any('note' in col for col in self.schema.columns) else 2
        for row_idx, row_data in enumerate(example_data, start_row):
            for col_idx, col in enumerate(self.schema.columns, 1):
                value = row_data.get(col['name'], '')
                ws.cell(row=row_idx, column=col_idx, value=value)
        
        # إضافة تعليمات
        instructions_row = start_row + len(example_data) + 2
        ws.cell(row=instructions_row, column=1, value="تعليمات:")
        ws.cell(row=instructions_row + 1, column=1, value="• الحقول باللون البرتقالي مطلوبة")
        ws.cell(row=instructions_row + 2, column=1, value="• يمكنك حذف صفوف الأمثلة وإضافة بياناتك")
        ws.cell(row=instructions_row + 3, column=1, value="• تأكد من صحة أسماء المشاريع والموردين المرجعية")
        
        # حفظ في الذاكرة
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        # إنشاء الاستجابة
        response = make_response(output.getvalue())
        response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        response.headers['Content-Disposition'] = f'attachment; filename={self.resource_name}_template.xlsx'
        
        return response
    
    def _fetch_data(self) -> List[Dict]:
        """جلب البيانات من قاعدة البيانات"""
        query = self.model_class.query
        
        # إضافة فلتر المشروع إذا لزم الأمر
        if self._needs_project_filter():
            query = query.filter_by(project_id=self.current_project_id)
        
        # جلب البيانات
        records = query.all()
        
        # تحويل لـ dictionaries
        data = []
        for record in records:
            row = {}
            for col in self.schema.columns:
                col_name = col['name']
                
                # معالجة العلاقات
                if col_name == 'project_code' and hasattr(record, 'project'):
                    row[col_name] = record.project.code if record.project else ''
                elif col_name == 'supplier_name' and hasattr(record, 'supplier'):
                    row[col_name] = record.supplier.name if record.supplier else ''
                elif hasattr(record, col_name):
                    value = getattr(record, col_name)
                    # معالجة القيم الخاصة
                    if value is None:
                        row[col_name] = ''
                    elif isinstance(value, (int, float)):
                        row[col_name] = value
                    else:
                        row[col_name] = str(value)
                else:
                    row[col_name] = ''
            
            data.append(row)
        
        return data
    
    def _needs_project_filter(self) -> bool:
        """التحقق من الحاجة لفلتر المشروع"""
        return hasattr(self.model_class, 'project_id') and self.current_project_id
    
    def _get_export_columns(self) -> List[str]:
        """الحصول على أسماء الأعمدة للتصدير"""
        return [col['name'] for col in self.schema.columns]