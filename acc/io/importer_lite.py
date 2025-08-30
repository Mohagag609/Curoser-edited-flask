"""
محرك الاستيراد الخفيف - بدون pandas
"""
import csv
import io
from typing import List, Dict, Any, Tuple, Optional
from flask import g
from sqlalchemy import and_, or_
from sqlalchemy.exc import IntegrityError
from acc.extensions import db
from acc.models import Project
from .schemas import get_schema, ResourceSchema

# محاولة استيراد pandas و openpyxl
try:
    import pandas as pd
    PANDAS_AVAILABLE = True
except ImportError:
    PANDAS_AVAILABLE = False

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

class ImportResult:
    """نتيجة عملية الاستيراد"""
    def __init__(self):
        self.inserted = 0
        self.updated = 0
        self.skipped_duplicates = 0
        self.errors: List[Dict[str, Any]] = []
        
    @property
    def total_processed(self):
        return self.inserted + self.updated + self.skipped_duplicates + len(self.errors)
    
    @property
    def success(self):
        return len(self.errors) == 0 or (self.inserted + self.updated > 0)
    
    def to_dict(self):
        return {
            'ok': self.success,
            'message': self.get_message(),
            'report': {
                'inserted': self.inserted,
                'updated': self.updated,
                'skipped_duplicates': self.skipped_duplicates,
                'errors': self.errors[:10]  # أول 10 أخطاء فقط
            }
        }
    
    def get_message(self):
        parts = []
        if self.inserted > 0:
            parts.append(f"تم الإدراج: {self.inserted}")
        if self.updated > 0:
            parts.append(f"تم التحديث: {self.updated}")
        if self.skipped_duplicates > 0:
            parts.append(f"مكرر: {self.skipped_duplicates}")
        if len(self.errors) > 0:
            parts.append(f"أخطاء: {len(self.errors)}")
        
        return " | ".join(parts) if parts else "لم يتم معالجة أي سجلات"

class GenericImporter:
    """مستورد عام للبيانات"""
    
    def __init__(self, resource_name: str):
        self.resource_name = resource_name
        self.schema = get_schema(resource_name)
        if not self.schema:
            raise ValueError(f"مورد غير معرف: {resource_name}")
        
        self.model_class = self.schema.model_class
        current_project = g.get('current_project')
        self.current_project_id = current_project.id if current_project and hasattr(current_project, 'id') else None
    
    def import_file(self, file_content: bytes, filename: str, mode: str = 'insert') -> ImportResult:
        """استيراد ملف"""
        result = ImportResult()
        
        # تحديد نوع الملف
        file_ext = filename.lower().split('.')[-1]
        
        try:
            if file_ext == 'csv':
                rows = self._read_csv_lite(file_content)
            elif file_ext in ['xlsx', 'xls']:
                if not OPENPYXL_AVAILABLE:
                    result.errors.append({
                        'row': 0,
                        'error': 'دعم Excel غير متوفر. يرجى استخدام CSV.'
                    })
                    return result
                rows = self._read_excel_lite(file_content)
            else:
                result.errors.append({
                    'row': 0,
                    'error': f'صيغة ملف غير مدعومة: {file_ext}'
                })
                return result
        except Exception as e:
            result.errors.append({
                'row': 0,
                'error': f'خطأ في قراءة الملف: {str(e)}'
            })
            return result
        
        # معالجة البيانات
        return self._process_rows(rows, mode, result)
    
    def _read_csv_lite(self, content: bytes) -> List[Dict[str, Any]]:
        """قراءة CSV بدون pandas"""
        # محاولة اكتشاف الترميز
        try:
            text = content.decode('utf-8-sig')
        except:
            try:
                text = content.decode('utf-8')
            except:
                text = content.decode('windows-1256')
        
        # قراءة CSV
        reader = csv.DictReader(io.StringIO(text))
        rows = []
        for row in reader:
            # تنظيف أسماء الأعمدة
            clean_row = {}
            for key, value in row.items():
                if key:
                    clean_key = key.strip().lower()
                    clean_row[clean_key] = value.strip() if value else ''
            rows.append(clean_row)
        
        return rows
    
    def _read_excel_lite(self, content: bytes) -> List[Dict[str, Any]]:
        """قراءة Excel بدون pandas"""
        from openpyxl import load_workbook
        
        # تحميل الملف
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        
        rows = []
        headers = []
        
        # قراءة العناوين من الصف الأول
        for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
            if row_idx == 0:
                headers = [str(h).strip().lower() if h else '' for h in row]
            else:
                # إنشاء dictionary من الصف
                row_dict = {}
                for col_idx, value in enumerate(row):
                    if col_idx < len(headers) and headers[col_idx]:
                        row_dict[headers[col_idx]] = str(value).strip() if value is not None else ''
                
                # تخطي الصفوف الفارغة
                if any(row_dict.values()):
                    rows.append(row_dict)
        
        wb.close()
        return rows
    
    def _process_rows(self, rows: List[Dict[str, Any]], mode: str, result: ImportResult) -> ImportResult:
        """معالجة الصفوف"""
        if not rows:
            result.errors.append({
                'row': 0,
                'error': 'الملف فارغ أو لا يحتوي على بيانات'
            })
            return result
        
        # التحقق من الأعمدة المطلوبة
        if rows:
            first_row_keys = set(rows[0].keys())
            missing_required = []
            for col in self.schema.required_columns:
                if col not in first_row_keys:
                    missing_required.append(self.schema.column_map[col]['label'])
            
            if missing_required:
                result.errors.append({
                    'row': 0,
                    'error': f'أعمدة مطلوبة مفقودة: {", ".join(missing_required)}'
                })
                return result
        
        # تحميل السجلات الموجودة للمقارنة
        existing_records = self._load_existing_records()
        
        # معالجة كل صف
        for idx, row in enumerate(rows):
            row_num = idx + 2  # رقم الصف في Excel (1-based + header)
            self._process_row(row, row_num, mode, existing_records, result)
        
        # حفظ التغييرات
        if result.success and (result.inserted > 0 or result.updated > 0):
            try:
                db.session.commit()
            except IntegrityError as e:
                db.session.rollback()
                result.errors.append({
                    'row': 0,
                    'error': 'تعارض في مفتاح فريد. تم إلغاء العملية.'
                })
                result.inserted = 0
                result.updated = 0
        
        return result
    
    def _load_existing_records(self) -> Dict[Tuple, Any]:
        """تحميل السجلات الموجودة"""
        query = self.model_class.query
        
        # إضافة فلتر المشروع إذا لزم الأمر
        if self._needs_project_filter():
            query = query.filter_by(project_id=self.current_project_id)
        
        records = {}
        for record in query.all():
            key = self._get_unique_key(record)
            if key:
                records[key] = record
        
        return records
    
    def _needs_project_filter(self) -> bool:
        """التحقق من الحاجة لفلتر المشروع"""
        return hasattr(self.model_class, 'project_id') and self.current_project_id
    
    def _get_unique_key(self, obj_or_row) -> Optional[Tuple]:
        """استخراج المفتاح الفريد"""
        values = []
        
        for field in self.schema.unique_by:
            if hasattr(obj_or_row, field):
                # كائن SQLAlchemy
                value = getattr(obj_or_row, field)
            elif isinstance(obj_or_row, dict):
                # صف من البيانات
                value = obj_or_row.get(field)
            else:
                return None
            
            # معالجة خاصة لـ project_id
            if field == 'project_id' and value is None:
                value = self.current_project_id
            
            values.append(value)
        
        return tuple(values) if all(v is not None for v in values) else None
    
    def _process_row(self, row: Dict[str, Any], row_num: int, mode: str, 
                     existing_records: Dict, result: ImportResult):
        """معالجة صف واحد"""
        # تطبيع البيانات
        normalized_data = {}
        errors = []
        
        for col_name, col_def in self.schema.column_map.items():
            if col_name in row and row[col_name]:
                value = row[col_name]
                
                # تطبيق دالة التطبيع
                if 'normalizer' in col_def:
                    try:
                        value = col_def['normalizer'](value)
                    except Exception as e:
                        errors.append(f"خطأ في تطبيع {col_def['label']}: {str(e)}")
                        continue
                
                # التحقق من القيم المسموحة
                if 'choices' in col_def and value and value not in col_def['choices']:
                    errors.append(f"{col_def['label']}: قيمة غير صالحة '{value}'")
                    continue
                
                normalized_data[col_name] = value
        
        # التحقق من الحقول المطلوبة
        for req_col in self.schema.required_columns:
            if req_col not in normalized_data or not normalized_data[req_col]:
                errors.append(f"{self.schema.column_map[req_col]['label']} مطلوب")
        
        if errors:
            result.errors.append({
                'row': row_num,
                'error': ' | '.join(errors)
            })
            return
        
        # معالجة العلاقات الخاصة
        self._process_relationships(normalized_data, row, errors)
        
        if errors:
            result.errors.append({
                'row': row_num,
                'error': ' | '.join(errors)
            })
            return
        
        # التحقق من التكرار
        unique_key = self._get_unique_key(normalized_data)
        existing = existing_records.get(unique_key) if unique_key else None
        
        if existing:
            if mode == 'insert':
                result.skipped_duplicates += 1
            else:  # upsert
                # تحديث الحقول غير الفارغة فقط
                updated = False
                for key, value in normalized_data.items():
                    if value and hasattr(existing, key):
                        old_value = getattr(existing, key)
                        if old_value != value:
                            setattr(existing, key, value)
                            updated = True
                
                if updated:
                    result.updated += 1
                else:
                    result.skipped_duplicates += 1
        else:
            # إنشاء سجل جديد
            try:
                new_obj = self.model_class(**normalized_data)
                
                # إضافة project_id إذا لزم الأمر
                if self._needs_project_filter():
                    new_obj.project_id = self.current_project_id
                
                # توليد الكود إذا لزم الأمر
                if hasattr(new_obj, 'generate_code') and not getattr(new_obj, 'code', None):
                    new_obj.generate_code()
                
                db.session.add(new_obj)
                result.inserted += 1
                
                # إضافة للسجلات الموجودة لتجنب التكرار في نفس الملف
                if unique_key:
                    existing_records[unique_key] = new_obj
                    
            except Exception as e:
                result.errors.append({
                    'row': row_num,
                    'error': f'خطأ في الإنشاء: {str(e)}'
                })
    
    def _process_relationships(self, data: Dict, row: Dict[str, Any], errors: List[str]):
        """معالجة العلاقات (مثل المشروع، المورد)"""
        # معالجة كود المشروع للوحدات
        if self.resource_name == 'units' and 'project_code' in row:
            project_code = row.get('project_code')
            if project_code:
                project = Project.query.filter_by(code=project_code).first()
                if project:
                    data['project_id'] = project.id
                    data.pop('project_code', None)
                else:
                    errors.append(f"كود المشروع '{project_code}' غير موجود")
        
        # معالجة اسم المورد للمواد
        elif self.resource_name == 'materials' and 'supplier_name' in row:
            supplier_name = row.get('supplier_name')
            if supplier_name:
                from acc.models import Supplier
                supplier = Supplier.query.filter_by(name=supplier_name).first()
                if supplier:
                    data['supplier_id'] = supplier.id
                    data.pop('supplier_name', None)
                else:
                    errors.append(f"اسم المورد '{supplier_name}' غير موجود")